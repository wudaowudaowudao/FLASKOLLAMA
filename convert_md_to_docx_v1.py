#新版本格式转换，表格内容还需要调整
import markdown2
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
import os


def set_font(run):
    run.font.name = '宋体'
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def md_to_docx(md_file_path, docx_file_path):
    # 读取Markdown文件
    with open(md_file_path, 'r', encoding='utf-8') as file:
        md_content = file.read()

    # 将Markdown内容转换为HTML
    html_content = markdown2.markdown(md_content)

    # 解析HTML
    soup = BeautifulSoup(html_content, 'html.parser')

    # 创建Word文档
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 遍历所有标签
    for elem in soup.find_all(['h1', 'h2', 'h3', 'p', 'ul', 'ol', 'img', 'table']):
        if elem.name == 'h1':
            para = doc.add_heading(elem.get_text(strip=True), level=1)
        elif elem.name == 'h2':
            para = doc.add_heading(elem.get_text(strip=True), level=2)
        elif elem.name == 'h3':
            para = doc.add_heading(elem.get_text(strip=True), level=3)
        elif elem.name == 'p':
            para = doc.add_paragraph(elem.get_text(strip=True))
        elif elem.name == 'ul':
            for li in elem.find_all('li'):
                para = doc.add_paragraph(li.get_text(strip=True), style='List Bullet')
        elif elem.name == 'ol':
            for li in elem.find_all('li'):
                para = doc.add_paragraph(li.get_text(strip=True), style='List Number')
        elif elem.name == 'img':
            src = elem.get('src')
            if src:
                img_name = os.path.basename(src)
                img_path = os.path.abspath(os.path.join('output/zhang1/images', img_name))
                if os.path.exists(img_path):
                    try:
                        para = doc.add_paragraph()
                        run = para.add_run()
                        run.add_picture(img_path, width=Inches(4))
                        set_font(run)
                        print(f"图片 {img_path} 插入成功")
                    except Exception as e:
                        print(f"图片 {img_path} 插入失败: {e}")
                else:
                    print(f"图片 {img_path} 不存在")
        elif elem.name == 'table':
            rows = elem.find_all('tr')
            if not rows:
                continue
            num_cols = len(rows[0].find_all(['th', 'td']))
            table = doc.add_table(rows=0, cols=num_cols)
            table.style = 'Table Grid'

            for row in rows:
                table_row = table.add_row()
                cells = row.find_all(['th', 'td'])
                for idx, cell in enumerate(cells):
                    if idx < num_cols:
                        table_row.cells[idx].text = cell.get_text(strip=True)

    # 保存文档
    doc.save(docx_file_path)
    print(f"成功保存到 {docx_file_path}")


if __name__ == "__main__":
    md_file_path = r'E:\PycharmProjects\FlaskOllama\output\zhang1\new\zhang1_zh.md'
    docx_file_path = r'E:\PycharmProjects\FlaskOllama\output\zhang1\new\zhang1_zh.docx'
    md_to_docx(md_file_path, docx_file_path)
