#图片导入仍有问题
import json
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from bs4 import BeautifulSoup

# 自定义路径变量
input_path = 'data/jsonToWord/Einleitung3.json'
output_path = 'output/jsonToWord/output.docx'

# 读取 JSON 文件
with open(input_path, 'r', encoding='utf-8') as f:
    data = json.load(f)

doc = Document()

# 设置全局字体为“宋体”
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
font.size = Pt(12)

def add_content(doc, item):
    if item['type'] == 'text':
        text = item.get('text', '').strip()
        level = item.get('text_level', None)
        if level == 1:
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)

    elif item['type'] == 'image':
        img_path = item.get('img_path', '')
        if os.path.exists(img_path):
            print(f"[图片导入成功] {img_path}")
            doc.add_picture(img_path, width=Inches(5))
        else:
            print(f"[图片未找到] {img_path}")
        if item.get('img_caption'):
            doc.add_paragraph("图注：" + ''.join(item['img_caption']))
        if item.get('img_footnote'):
            doc.add_paragraph("脚注：" + ''.join(item['img_footnote']))

    elif item['type'] == 'table':
        html = item.get('table_body', '')
        soup = BeautifulSoup(html, 'html.parser')
        table_tag = soup.find('table')
        if table_tag:
            rows = table_tag.find_all('tr')
            if rows:
                max_cols = max(len(row.find_all(['td', 'th'])) for row in rows)
                table = doc.add_table(rows=len(rows), cols=max_cols)
                table.style = 'Table Grid'
                for i, row in enumerate(rows):
                    cells = row.find_all(['td', 'th'])
                    for j, cell in enumerate(cells):
                        table.cell(i, j).text = cell.get_text(strip=True)
        else:
            print("[表格解析失败] 无法从 table_body 中提取 <table> 标签")

        if item.get('table_caption'):
            doc.add_paragraph("表注：" + ''.join(item['table_caption']))
        if item.get('table_footnote'):
            doc.add_paragraph("脚注：" + ''.join(item['table_footnote']))

# 添加内容
for item in data:
    add_content(doc, item)

# 保存 DOCX 文件
doc.save(output_path)
print(f"[完成] 文档已保存为：{output_path}")
