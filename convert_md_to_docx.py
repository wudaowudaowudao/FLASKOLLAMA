#表格生成估计有问题，电脑打不开，手机能打开

import markdown
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
import os

def md_to_docx(md_file_path, docx_file_path):
    # 读取Markdown文件内容
    with open(md_file_path, 'r', encoding='utf-8') as file:
        md_content = file.read()
    
    # 将Markdown内容转换为HTML
    html_content = markdown.markdown(md_content)
    
    # 使用BeautifulSoup解析HTML内容
    soup = BeautifulSoup(html_content, 'html.parser')
    
    # 创建一个新的Word文档
    doc = Document()
    
    # 定义设置字体的函数
    def set_font(run):
        run.font.name = '宋体'
        r = run._element
        r.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 遍历HTML元素并添加到Word文档中
    for element in soup.descendants:
        if element.name == 'h1':
            heading = doc.add_heading(element.get_text(), level=1)
            for run in heading.runs:
                set_font(run)
        elif element.name == 'h2':
            heading = doc.add_heading(element.get_text(), level=2)
            for run in heading.runs:
                set_font(run)
        elif element.name == 'h3':
            heading = doc.add_heading(element.get_text(), level=3)
            for run in heading.runs:
                set_font(run)
        elif element.name == 'p':
            paragraph = doc.add_paragraph(element.get_text())
            for run in paragraph.runs:
                set_font(run)
        elif element.name == 'ul':
            for li in element.find_all('li'):
                paragraph = doc.add_paragraph(li.get_text(), style='List Bullet')
                for run in paragraph.runs:
                    set_font(run)
        elif element.name == 'ol':
            for li in element.find_all('li'):
                paragraph = doc.add_paragraph(li.get_text(), style='List Number')
                for run in paragraph.runs:
                    set_font(run)
        elif element.name == 'img':
            # 提取图片文件名
            img_src = element.get('src')
            img_name = os.path.basename(img_src)
            local_img_path = os.path.join('output/zhang1/images', img_name)  # 假设图片在本地的 images 文件夹下
            absolute_img_path = os.path.abspath(local_img_path)

            if os.path.exists(absolute_img_path):
                try:
                    # 插入图片到 Word 文档
                    paragraph = doc.add_paragraph()
                    run = paragraph.add_run()
                    run.add_picture(absolute_img_path, width=Inches(4))  # 图片宽度设置为 4 英寸
                    set_font(run)
                    print(f"图片 {absolute_img_path} 插入成功")
                except Exception as e:
                    print(f"插入图片 {absolute_img_path} 时出错: {e}")
            else:
                print(f"图片 {absolute_img_path} 不存在")
        elif element.name == 'table':
            rows = element.find_all('tr')
            if not rows:
                continue

            # 创建表格，列数取第一行的单元格数
            num_cols = len(rows[0].find_all(['th', 'td']))
            table = doc.add_table(rows=0, cols=num_cols)

            # 设置表格样式
            table.style = 'Table Grid'  # 使用Word内置样式防止出错

            for row in rows:
                table_row = table.add_row()
                cells = row.find_all(['th', 'td'])
                for idx, cell in enumerate(cells):
                    if idx < len(table_row.cells):
                        table_row.cells[idx].text = cell.get_text()
                        for paragraph in table_row.cells[idx].paragraphs:
                            for run in paragraph.runs:
                                set_font(run)

            # 设置表格边框
            from docx.oxml import OxmlElement

            def set_table_border(table):
                tbl = table._tbl
                tblPr = tbl.tblPr  # 直接访问 tblPr
                if tblPr is None:
                    tblPr = OxmlElement('w:tblPr')
                    tbl.append(tblPr)

                # 检查是否已经存在表格边框设置
                tblBorders = tblPr.find(qn('w:tblBorders'))
                if tblBorders is None:
                    tblBorders = OxmlElement('w:tblBorders')
                    tblPr.append(tblBorders)

                border_style = {
                    'val': 'single',
                    'sz': '4',
                    'space': '0',
                    'color': 'auto'
                }

                for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                    border = OxmlElement(f'w:{border_name}')
                    for key, value in border_style.items():
                        border.set(qn(f'w:{key}'), value)
                    tblBorders.append(border)

            set_table_border(table)

    # 保存Word文档
    doc.save(docx_file_path)

if __name__ == "__main__":
    md_file_path = r'E:\PycharmProjects\FlaskOllama\output\zhang1\zhang1_zh.md'
    docx_file_path = r'output/zhang1_zh.docx'
    md_to_docx(md_file_path, docx_file_path)