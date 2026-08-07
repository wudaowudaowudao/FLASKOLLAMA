#初步转换，未考虑表格
import json
from docx import Document
from docx.shared import Inches
import os

# 读取 JSON 文件
with open('E:\PycharmProjects\FlaskOllama\data\jsonToWord\Einleitung3.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

doc = Document()

def add_content(doc, item):
    if item['type'] == 'text':
        text = item['text']
        if 'text_level' in item and item['text_level'] == 1:
            doc.add_heading(text.strip(), level=1)
        else:
            doc.add_paragraph(text.strip())
    elif item['type'] == 'image':
        if os.path.exists(item['img_path']):
            doc.add_picture(item['img_path'], width=Inches(5))
        if item.get('img_caption'):
            doc.add_paragraph("图注：" + ''.join(item['img_caption']))
        if item.get('img_footnote'):
            doc.add_paragraph("脚注：" + ''.join(item['img_footnote']))
    elif item['type'] == 'table':
        doc.add_paragraph("表格请参考下图：")
        if os.path.exists(item['img_path']):
            doc.add_picture(item['img_path'], width=Inches(5))
        if item.get('table_caption'):
            doc.add_paragraph("表注：" + ''.join(item['table_caption']))
        if item.get('table_footnote'):
            doc.add_paragraph("脚注：" + ''.join(item['table_footnote']))

for item in data:
    add_content(doc, item)

doc.save('output/output.docx')
print("保存为 output.docx")
