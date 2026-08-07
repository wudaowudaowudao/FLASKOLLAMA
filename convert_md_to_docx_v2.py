#用于md转docx，已经解决公式转换的问题
import markdown2
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from docx.enum.text import WD_BREAK
import os
import re
import latex2mathml.converter
from lxml import etree

# 匹配 $...$ 中的 LaTeX 公式
LATEX_MATH_PATTERN = re.compile(r'\$(.+?)\$')

# 设置字体为宋体
def set_font(run):
    run.font.name = '宋体'
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# 将 latex 公式转为 OMML（Word 数学格式）
def latex_to_omml(latex_expr):
    try:
        # Step 1: LaTeX to MathML
        mathml_str = latex2mathml.converter.convert(latex_expr)
        mathml_xml = etree.fromstring(mathml_str.encode('utf-8'))

        # Step 2: MathML to OMML
        xslt_path = os.path.join(os.path.dirname(__file__), 'mml2omml.xsl')
        xslt_doc = etree.parse(xslt_path)
        transform = etree.XSLT(xslt_doc)
        omml_tree = transform(mathml_xml)
        omml_str = etree.tostring(omml_tree, encoding='unicode')
        omml_element = parse_xml(omml_str)
        return omml_element
    except Exception as e:
        print(f"[LaTeX 转换失败] {latex_expr} 错误: {e}")
        return None

# 替换段落中的 LaTeX 为 Word 数学公式
def replace_latex_with_omml(paragraph, text):
    parts = LATEX_MATH_PATTERN.split(text)
    for i, part in enumerate(parts):
        if i % 2 == 0:
            if part:
                run = paragraph.add_run(part)
                set_font(run)
        else:
            omml = latex_to_omml(part)
            if omml is not None:
                paragraph._element.append(omml)
            else:
                run = paragraph.add_run(f"${part}$")
                set_font(run)

# Markdown 转 Word 主逻辑
def md_to_docx(md_file_path, docx_file_path):
    with open(md_file_path, 'r', encoding='utf-8') as file:
        md_content = file.read()

    html_content = markdown2.markdown(md_content)
    soup = BeautifulSoup(html_content, 'html.parser')

    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    for elem in soup.find_all(['h1', 'h2', 'h3', 'p', 'ul', 'ol', 'img', 'table']):
        if elem.name == 'h1':
            doc.add_heading(elem.get_text(strip=True), level=1)
        elif elem.name == 'h2':
            doc.add_heading(elem.get_text(strip=True), level=2)
        elif elem.name == 'h3':
            doc.add_heading(elem.get_text(strip=True), level=3)
        elif elem.name == 'p':
            para = doc.add_paragraph()
            replace_latex_with_omml(para, elem.get_text())
        elif elem.name == 'ul':
            for li in elem.find_all('li'):
                para = doc.add_paragraph(li.get_text(strip=True), style='List Bullet')
                set_font(para.runs[0])
        elif elem.name == 'ol':
            for li in elem.find_all('li'):
                para = doc.add_paragraph(li.get_text(strip=True), style='List Number')
                set_font(para.runs[0])
        elif elem.name == 'img':
            src = elem.get('src')
            if src:
                img_path = os.path.abspath(src)
                if os.path.exists(img_path):
                    para = doc.add_paragraph()
                    run = para.add_run()
                    run.add_picture(img_path, width=Inches(4))
                    set_font(run)
                else:
                    print(f"图片不存在: {img_path}")
        elif elem.name == 'table':
            rows = elem.find_all('tr')
            if not rows:
                continue
            num_cols = len(rows[0].find_all(['th', 'td']))
            table = doc.add_table(rows=0, cols=num_cols)
            table.style = 'Table Grid'
            for row in rows:
                table_row = table.add_row()
                cells = row.find_all(['th', 'td'])
                for idx, cell in enumerate(cells):
                    if idx < num_cols:
                        table_row.cells[idx].text = cell.get_text(strip=True)

    doc.save(docx_file_path)
    print(f"✅ Word 文档已保存到: {docx_file_path}")

# 示例调用
if __name__ == "__main__":
    md_file_path = r"E:\PycharmProjects\FlaskOllama\output\zhang1\new\zhang1_zh.md"
    docx_file_path = r"E:\PycharmProjects\FlaskOllama\output\zhang1\new\zhang1_zh.docx"
    md_to_docx(md_file_path, docx_file_path)
