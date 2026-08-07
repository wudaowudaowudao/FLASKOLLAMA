import docx
import json
import os
from typing import Optional, Dict, List, Union

class DocxToJsonConverter:
    def __init__(self, docx_file: str):
        """初始化转换器

        Args:
            docx_file (str): Word文档的路径
        """
        self.docx_file = docx_file
        self.data: Optional[Dict] = None

    def read_docx(self) -> 'DocxToJsonConverter':
        """读取Word文档

        Returns:
            self: 支持链式调用
        """
        doc = docx.Document(self.docx_file)
        
        # 提取文档结构和内容
        self.data = {
            'paragraphs': [],
            'tables': []
        }
        
        # 提取段落
        for para in doc.paragraphs:
            if para.text.strip():
                self.data['paragraphs'].append({
                    'text': para.text,
                    'style': para.style.name if para.style else 'Normal'
                })
        
        # 提取表格
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    row_data.append(cell.text.strip())
                table_data.append(row_data)
            self.data['tables'].append(table_data)
        
        return self

    def save_json(self, output_file: Optional[str] = None, indent: int = 4) -> str:
        """将数据保存为JSON文件

        Args:
            output_file: 输出JSON文件的路径，如果为None则使用Word文档同名的JSON文件
            indent: JSON文件的缩进空格数

        Returns:
            str: 保存的JSON文件路径
        """
        if self.data is None:
            raise ValueError('请先调用read_docx()方法读取Word文档')

        if output_file is None:
            output_file = os.path.splitext(self.docx_file)[0] + '.json'

        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(self.data, f, ensure_ascii=False, indent=indent)

        return output_file

    def get_data(self) -> Optional[Dict]:
        """获取转换后的数据

        Returns:
            Optional[Dict]: 转换后的数据，如果未读取文档则返回None
        """
        return self.data

if __name__ == '__main__':
    # 测试代码
    docx_file = 'e:\\PycharmProjects\\TestEnv\\XMLGenerator\\data\\合作报价及合同审计问题注意事项1011.docx'
    
    # 基本使用示例
    converter = DocxToJsonConverter(docx_file)
    json_file = converter.read_docx().save_json()
    print(f'已将Word文档转换为JSON并保存至: {json_file}')

    # 获取数据示例
    data = converter.get_data()
    if data:
        print(f'\n文档结构预览:')
        print(f'段落数量: {len(data["paragraphs"])}')
        print(f'表格数量: {len(data["tables"])}')
        
        if data['paragraphs']:
            print(f'\n第一段内容预览: {data["paragraphs"][0]["text"][:100]}...')

    # 自定义输出文件示例
    custom_output = 'e:\\PycharmProjects\\TestEnv\\XMLGenerator\\output\\custom_output.json'
    json_file = converter.save_json(custom_output)
    print(f'\n已将数据保存至自定义路径: {json_file}')