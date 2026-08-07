import os
import time
import random
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import OllamaEmbeddings
import docx


class WordFaissCreator:
    def __init__(self, model="bge-m3:latest", show_progress=True):
        self.embeddings = OllamaEmbeddings(model=model, show_progress=show_progress)

    def load_word_data(self, file_path):
        """
        读取 Word 文件内容（包括段落和表格）
        """
        doc = docx.Document(file_path)
        full_text = []

        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(' | '.join(row_text))

        return '\n'.join(full_text)

    def split_text(self, text, chunk_size=10, overlap=2):
        paragraphs = [p.strip() for p in text.split('\n') if p.strip()]
        chunks = []
        i = 0
        while i < len(paragraphs):
            end = i + chunk_size
            chunk = '\n'.join(paragraphs[i:end])
            chunks.append(chunk)
            i += chunk_size - overlap
        return chunks

    def process_and_embed_chunks(self, chunks, batch_size=1, initial_delay=0.0, max_retries=5):
        all_texts = []
        for i in range(0, len(chunks), batch_size):
            batch = chunks[i:i + batch_size]
            retries = 0
            delay = initial_delay

            while retries < max_retries:
                try:
                    _ = self.embeddings.embed_documents(batch)
                    all_texts.extend(batch)
                    break
                except Exception as e:
                    retries += 1
                    print(f"Error: {e}")
                    if retries < max_retries:
                        wait = delay * (2 ** retries) + random.uniform(0, 0.1)
                        print(f"Retrying in {wait:.2f}s...")
                        time.sleep(wait)
                    else:
                        print("Max retries reached for this batch.")
        return all_texts

    def create_faiss(self, word_file_path, save_path):
        word_text = self.load_word_data(word_file_path)
        print('Loaded word text, length:', len(word_text))

        chunks = self.split_text(word_text)
        print('Total chunks created:', len(chunks))

        all_texts = self.process_and_embed_chunks(chunks)
        print('Total embedded chunks:', len(all_texts))

        db = FAISS.from_texts(all_texts, self.embeddings)

        file_name = os.path.splitext(os.path.basename(word_file_path))[0]
        db.save_local(save_path, index_name=file_name)
        print(f'FAISS vector store saved at {save_path}')
        return save_path


# 主函数
def main():
    word_file_path = r'E:\PycharmProjects\FlaskOllama\data\SimuWord.docx'
    save_path = r"E:\PycharmProjects\FlaskOllama\data\vector_word_db"

    creator = WordFaissCreator()
    creator.create_faiss(word_file_path, save_path)


if __name__ == "__main__":
    main()
