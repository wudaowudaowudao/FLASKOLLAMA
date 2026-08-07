# xml_rule_checker.py
import os
import json
import platform
import xml.etree.ElementTree as ET
from xml.dom import minidom
from langchain import PromptTemplate
from langchain_community.chat_models import ChatOllama
from docx import Document


class XMLRuleChecker:
    # 默认LLM配置
    DEFAULT_LLM_CONFIG = {
        'type': 'zhipu',
        'api_key': 'd6366483a4ae45fb9a4be838eb222ee7.jr7A8GlghBmdbkgM',
        'model': 'glm-4-air-250414',
        'temperature': 0.1
    }

    FORMAT_JSON_NEWLINES = True

    @staticmethod
    def convert_xml_to_str(file):
        with open(file, 'r', encoding='utf-8') as f:
            return f.read()

    @staticmethod
    def convert_docx_to_xml(docx_file_path):
        try:
            doc = Document(docx_file_path)
            root = ET.Element("Document")

            doc_info = ET.SubElement(root, "DocumentInfo")
            ET.SubElement(doc_info, "SourceFile").text = os.path.basename(docx_file_path)
            ET.SubElement(doc_info, "ConvertedFrom").text = "DOCX"

            paragraphs_elem = ET.SubElement(root, "Paragraphs")
            for i, paragraph in enumerate(doc.paragraphs):
                if paragraph.text.strip():
                    para_elem = ET.SubElement(paragraphs_elem, "Paragraph")
                    para_elem.set("id", str(i + 1))
                    para_elem.text = paragraph.text.strip()

            tables_elem = ET.SubElement(root, "Tables")
            for table_idx, table in enumerate(doc.tables):
                table_elem = ET.SubElement(tables_elem, "Table")
                table_elem.set("id", str(table_idx + 1))
                for row_idx, row in enumerate(table.rows):
                    row_elem = ET.SubElement(table_elem, "Row")
                    row_elem.set("id", str(row_idx + 1))
                    for cell_idx, cell in enumerate(row.cells):
                        cell_elem = ET.SubElement(row_elem, "Cell")
                        cell_elem.set("id", str(cell_idx + 1))
                        cell_elem.text = cell.text.strip() if cell.text else ""

            xml_str = ET.tostring(root, encoding='unicode', method='xml')
            dom = minidom.parseString(xml_str)
            formatted_xml = dom.toprettyxml(indent="  ", encoding=None)
            return '\n'.join([line for line in formatted_xml.split('\n') if line.strip()])
        except Exception as e:
            print(f"转换DOCX文件时出错: {e}")
            return None

    def get_target_file_content(self, target_file):
        ext = os.path.splitext(target_file)[1].lower()
        if ext == '.xml':
            return self.convert_xml_to_str(target_file)
        elif ext == '.docx':
            return self.convert_docx_to_xml(target_file)
        else:
            raise Exception(f"不支持的文件格式: {ext}")

    def process_single_pair(self, rule_str, target_str):
        llm_config = self.llm_config or self.DEFAULT_LLM_CONFIG.copy()
        if llm_config == self.DEFAULT_LLM_CONFIG:
            print(">>> 使用默认LLM配置")

        template = """rule_file是xml格式的规则文件内容，其中，RuleItem是单个规则项，Description是规则逻辑描述，
CheckConfig下的Rule是规则逻辑涉及的参数。
result_file是从模型中提取的信息。
检查rule_file中的每一条规则，与result_file获取到的相关的内容进行检查，是否一致
rule_file: {xml1}
result_file: {xml2}

Question: 逐条检查rule_file中的规则，与result_file获取到的相关的内容进行检查，写出rule_file每一条RuleItem的对比结果，回答用中文

Result:"""

        prompt = PromptTemplate(
            input_variables=["xml1", "xml2"],
            template=template
        )

        if llm_config.get('type') == 'zhipu':
            try:
                from langchain_community.chat_models import ChatZhipuAI
                llm = ChatZhipuAI(
                    model=llm_config['model'],
                    temperature=llm_config.get('temperature', 0.1),
                    zhipuai_api_key=llm_config['api_key'],
                    request_timeout=1200,
                    timeout=1200
                )
            except Exception as e:
                print(f"ZhipuAI 初始化失败：{e}，回退到本地ollama")
                llm_config['type'] = 'ollama'

        if llm_config.get('type') == 'ollama':
            os_type = platform.system().lower()
            model_name = "qwen3:235b" if os_type == "linux" else "gemma3:27b"
            print(f">>> 使用本地Ollama模型：{model_name}")
            llm = ChatOllama(model=model_name, temperature=0.1, num_ctx=48 * 1000, num_predict=64 * 1000)

        chain = prompt | llm

        for attempt in range(20):
            try:
                print(f"第 {attempt + 1} 次尝试调用LLM...")
                rs = chain.invoke({"xml1": rule_str, "xml2": target_str})
                return rs.content if hasattr(rs, 'content') else str(rs)
            except Exception as e:
                print(f"LLM尝试失败（第{attempt + 1}次），等待重试... {e}")
                import time
                time.sleep(5)
        raise RuntimeError("所有LLM重试失败")

    def __init__(self, rule_file, target_file, output_folder="output", llm_config=None):
        # 支持单个规则文件路径或规则文件路径列表
        if isinstance(rule_file, str):
            self.rule_files = [rule_file]
        elif isinstance(rule_file, list) and all(isinstance(f, str) for f in rule_file):
            self.rule_files = rule_file
        else:
            raise TypeError("rule_file参数必须是字符串或字符串列表")

        self.target_file = target_file
        self.output_folder = output_folder
        self.llm_config = llm_config
        os.makedirs(self.output_folder, exist_ok=True)

    def run(self):
        try:
            # 读取所有规则文件并合并内容
            rule_strs = []
            for rule_file in self.rule_files:
                rule_strs.append(self.convert_xml_to_str(rule_file))
            combined_rule_str = "\n\n".join(rule_strs)

            target_str = self.get_target_file_content(self.target_file)
            result = self.process_single_pair(combined_rule_str, target_str)

            name_without_ext = os.path.splitext(os.path.basename(self.target_file))[0]
            output_filename = f"Check_output_{name_without_ext}.json"
            output_path = os.path.join(self.output_folder, output_filename)

            output_data = {
                "rule_files": [os.path.basename(f) for f in self.rule_files],
                "target_file": os.path.basename(self.target_file),
                "target_file_format": os.path.splitext(self.target_file)[1].lower(),
                "processing_time": "",
                "result": result
            }

            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(output_data, f, ensure_ascii=False, indent=2)

            if self.FORMAT_JSON_NEWLINES:
                self._format_newlines(output_path)

            print(f"\n✅ 审查完成，结果保存在：{output_path}")
            return output_data

        except Exception as e:
            print(f"❌ 处理文件失败: {e}")
            return None

    def _format_newlines(self, file_path):
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        formatted_content = content.replace('\\n', '\n')
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(formatted_content)


# ==================== 示例入口 ====================
if __name__ == "__main__":
    # 示例路径，可根据需要替换
    rule_path = "rule_files/simulated_rule_CAH234_03_00_001_A_dwg1.xml"
    target_path = "model_files/CTH11_01_04_000_A_dwg1.xml"
    output_dir = "output"

    # 示例：使用单个规则文件
    # checker = XMLRuleChecker(rule_file=rule_path, target_file=target_path, output_folder=output_dir)

    # 示例：使用多个规则文件
    rule_paths = [
        "rule_files/simulated_rule_CAH234_03_00_001_A_dwg1.xml",
        "rule_files/another_rule.xml"  # 替换为实际的规则文件路径
    ]
    checker = XMLRuleChecker(rule_file=rule_paths, target_file=target_path, output_folder=output_dir)
    result = checker.run()

    if result:
        print("\n🎯 审查结果内容预览：")
        print(json.dumps(result["result"], indent=2, ensure_ascii=False))
